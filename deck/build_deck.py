#!/usr/bin/env python3
"""Build the SolarGuard rubric deck (16:9 pptx) + an HTML preview of the same slides.

Design is lifted from the team's scroll-story page: near-black space background,
hairline cards, one warm accent (sun), one cool accent (orbit), mono labels.

    python3 deck/build_deck.py            # writes deck/SolarGuard-Rubric.pptx + preview.html
"""
from __future__ import annotations

import html
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUT = os.path.dirname(os.path.abspath(__file__))

# palette, from the scroll story
SPACE = RGBColor(0x04, 0x06, 0x0C)
CARD = RGBColor(0x0B, 0x10, 0x1B)
CARD2 = RGBColor(0x14, 0x1B, 0x29)
LINE = RGBColor(0x23, 0x2C, 0x3C)
INK = RGBColor(0xF3, 0xF4, 0xF6)
MUTED = RGBColor(0x9A, 0xA3, 0xB2)
SUN = RGBColor(0xFF, 0x8A, 0x2B)
SUNINK = RGBColor(0xFF, 0xA4, 0x5C)
ORBIT = RGBColor(0x5B, 0x7C, 0xFF)
CLEAN = RGBColor(0x4A, 0xDE, 0x9B)
DARKINK = RGBColor(0x0E, 0x11, 0x17)

HEAD = "Arial"
BODY = "Calibri"
MONO = "Courier New"

W, H = 13.333, 7.5
PAL = {"space": "#04060c", "card": "#0b101b", "card2": "#141b29", "line": "#232c3c",
       "ink": "#f3f4f6", "muted": "#9aa3b2", "sun": "#ff8a2b", "sunink": "#ffa45c",
       "orbit": "#5b7cff", "clean": "#4ade9b", "darkink": "#0e1117"}

# ───────────────────────────────────────────────────────── the content
SLIDES: list[dict] = [
    {"kind": "title",
     "eyebrow": "Dust intelligence for Saudi solar",
     "big": "SolarGuard",
     "sub": "You can't see the dust. We can.",
     "foot": "Team space-Marines  ·  Riyadh\u2013Dammam",
     "notes": "Open cold. One sentence: SolarGuard tells a Saudi solar plant, every morning, whether "
              "cleaning the panels today is worth the money. Then move."},

    {"kind": "cards3",
     "eyebrow": "What it is",
     "title": "Satellite in. One decision out.",
     "cards": [("Watch", "12 live open feeds: satellite dust, aerosol, wind, sun, rain.", "12 feeds"),
               ("Predict", "Our own AI models, trained in-house on 4 years of Saudi data.", "5 models"),
               ("Decide", "CLEAN or WAIT, priced in riyals. Nothing installed on site.", "1 verdict")],
     "notes": "Three verbs. Do not explain the internals yet, the rubric slides do that."},

    {"kind": "pipeline",
     "eyebrow": "How it works",
     "title": "Five steps, satellite to decision",
     "steps": [("1", "Live data"), ("2", "AI models"), ("3", "Forecasts"),
               ("4", "Economics"), ("5", "CLEAN or WAIT")],
     "note": "The AI predicts. The physics engine decides.",
     "notes": "Emphasise step 5: the output is a decision, not a chart."},

    {"kind": "rubric",
     "eyebrow": "The rubric",
     "title": "Why this is a 5 on all seven",
     "foot": "Each criterion, the requirement, and the evidence we can show.",
     "notes": "One slide to frame the next seven. Point out that every criterion gets its own slide."},

    {"kind": "criterion", "n": 1, "name": "Problem definition and specificity", "weight": "15%",
     "rows": [("a", "Dust costs 20-40% of output: measured on the Saudi east coast at 30-35%, and 20% from a single sandstorm.", "KAUST 2023 · Energies 2022"),
              ("b", "Today: fixed calendars (weekly, 10-day) or a dust threshold. None ask what tomorrow brings.", "IEA-PVPS 2022"),
              ("c", "At 100 MWp: 10.78M SAR a year cleaning weekly, 16.13M never cleaning.", "our simulation, measured weather")],
     "notes": "Sources are on the slide. If asked for the weakest number, say the 10.78/16.13 pair is "
              "our own simulation and name the assumptions (100 MWp, Dammam, 1,800 SAR/MWp per pass)."},

    {"kind": "criterion", "n": 2, "name": "Originality and differentiation", "weight": "15%",
     "rows": [("a", "Compared with the fixed calendar, and with dust-threshold sensors that need hardware and act after the loss.", "closest available approaches"),
              ("b", "We move cleaning from a calendar to a forecast plus economics: a riyal decision, no hardware.", "our contribution"),
              ("c", "Output error 0.234 vs 2.572 kWh/kWp for the naive guess: 11x. Trade-off: a brand-new site is less accurate (AUC 0.785).", "tested, limitation stated")],
     "notes": "Naming the trade-off is what the rubric asks for at 5, not a weakness to avoid."},

    {"kind": "criterion", "n": 3, "name": "Solution design and expected effectiveness", "weight": "20%",
     "rows": [("a", "Success criterion: beat the 10-day habit on annual cost at equal or better energy. Model: 6.71M vs 8.41M SAR a year, 18 cleans.", "measurable target"),
              ("b", "Failures handled: a dead feed degrades one number; if the AI is unavailable the physics engine answers; no input returns a clear error.", "safeguards, by design"),
              ("c", "Boundary of use: dust-dominated, low-rain desert sites. Snow, hydro and coastal fog are out of scope.", "stated, not implied")],
     "notes": "This criterion gives 0 without AI or ML. Say it plainly: five models we trained, and an "
              "agentic copilot."},

    {"kind": "criterion", "n": 4, "name": "Technical and implementation feasibility", "weight": "15%",
     "rows": [("a", "Core assumption (soiling 0.2-0.8% per day in Saudi) comes from published measurement, not guesswork.", "IEA-PVPS 2022 · KAUST 2020"),
              ("b", "Pilot: one 100 MWp site, 90 days, software only. Estimate: one engineer plus cloud, about 15k SAR a month.", "estimate, assumptions stated"),
              ("c", "Two real risks: no inverter ground truth, and feed outages. Mitigations: compare against the plant's own inverters; cache plus physics fallback.", "each with a fix")],
     "notes": "Flag the 15k SAR as our estimate, not a quote. If pressed, the cost is dominated by people, "
              "not infrastructure."},

    {"kind": "criterion", "n": 5, "name": "Demonstrated development and validation", "weight": "10%",
     "rows": [("a", "Working end to end: pick a site on the map, get the AI prediction, the verdict, and an agent answer with its tool trace.", "live, not mocked"),
              ("b", "Tested against baselines on 2,244 unseen days: storm AUC 0.952, output within 7.7%, soiling within 8.4%.", "measured"),
              ("c", "Failure tested: hide a whole site and score it (AUC 0.785), plus a test that fails if the served data drifts from training.", "reliability reported")],
     "notes": "If judges want the failure test live, the leave-one-site-out block is in the model metrics file."},

    {"kind": "criterion", "n": 6, "name": "Saudi-market relevance and potential impact", "weight": "15%",
     "rows": [("a", "Built on 12 Saudi sites from NEOM to Dammam, on Saudi dust literature, for the Kingdom's solar build-out.", "local evidence"),
              ("b", "Per 100 MWp site: about 1.7M SAR and 4.5M litres of water saved a year versus the 10-day habit.", "simulated, assumptions stated"),
              ("c", "Adopted by the plant's O&M manager: open the dashboard in the morning, read CLEAN or WAIT, send crews only when it pays.", "fits the workflow")],
     "notes": "Say the water number matters because every pass skipped is water not used in a desert."},

    {"kind": "criterion", "n": 7, "name": "Presentation and responses to questions", "weight": "10%",
     "rows": [("a", "One slide per criterion, in order, inside the time.", "structure"),
              ("b", "We label every number: built and tested, or estimated.", "no blurring"),
              ("c", "Answers come with the source, and we name what we don't know: at one day ahead the naive rule still wins on F1.", "uncertainties named")],
     "notes": "Then the extras slides. Keep this one short, it is about how we present."},

    {"kind": "stats",
     "eyebrow": "Extra · the models we trained",
     "title": "Five models, built here",
     "stats": [("5", "models shipped"), ("37", "inputs each"), ("1,540", "boosted trees")],
     "rows": ["Trained in-house on 4 years of Saudi dust and weather, 12 sites, 18,348 site-days.",
              "Each model blends a gradient-boosted tree ensemble with a neural network.",
              "Tested on 2,244 days never seen: storm AUC 0.952, output within 7.7%, soiling within 8.4%."],
     "notes": "Each of the five is two trained models blended, so ten in total. No GPU, no ML library: "
              "hand-written in numpy."},

    {"kind": "steps",
     "eyebrow": "Extra · the agentic AI",
     "title": "It asks, waits, then answers",
     "steps": [("1", "Operator asks a question"),
               ("2", "Model calls its tools"),
               ("3", "Server runs them and waits"),
               ("4", "Model answers with citations")],
     "rows": ["10 tools over the live data, plus document search so answers carry sources.",
              "Retrieval over 672 indexed chunks of soiling research.",
              "It acts on the page: moves the map, switches the satellite layer."],
     "notes": "The model never speaks from memory: it asks for data, waits for the real output, then talks."},

    {"kind": "close",
     "big": "CLEAN or WAIT",
     "sub": "From a fixed schedule to a data-driven economic decision.",
     "foot": "github.com/Hermes-mahers-bot/solarguard-space",
     "notes": "Stop talking. Let the last line land, then take questions."},
]


# ───────────────────────────────────────────────────────── pptx helpers
def slide_base(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = SPACE
    return s


def box(s, x, y, w, h, *, size=15, color=INK, bold=False, font=BODY, align=PP_ALIGN.LEFT,
        line_spacing=1.15, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    r = p.add_run()
    r.text = ""
    f = r.font
    f.size, f.bold, f.name = Pt(size), bold, font
    f.color.rgb = color
    return tb, r


def put(s, x, y, w, h, text, **kw):
    tb, r = box(s, x, y, w, h, **kw)
    r.text = text
    return tb


def card(s, x, y, w, h, *, fill=CARD, line=LINE, radius=0.055):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = radius
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    return sh


def chip(s, x, y, w, h, text, *, fill=CARD2, color=INK, size=11, font=MONO, bold=True):
    sh = card(s, x, y, w, h, fill=fill, line=None, radius=0.28)
    tf = sh.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.name = Pt(size), bold, font
    r.font.color.rgb = color
    return sh


def cols(n, gap=0.22, left=0.9, right=0.9):
    """Card width and pitch for n columns inside the margins, so a row can never
    run off the slide edge (the 5-step pipeline did when the width was fixed)."""
    total = W - left - right
    w = (total - (n - 1) * gap) / n
    return w, w + gap


def eyebrow(s, text, y=0.62):
    put(s, 0.9, y, 8.0, 0.3, text.upper(), size=11, color=SUNINK, bold=True, font=MONO)


def title(s, text, y=1.0, size=34, w=11.5):
    put(s, 0.9, y, w, 0.9, text, size=size, color=INK, bold=True, font=HEAD, line_spacing=1.02)


def no(story, slide):
    slide.notes_slide.notes_text_frame.text = story.get("notes", "")


def build_pptx(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)

    for i, sl in enumerate(SLIDES):
        k = sl["kind"]
        s = slide_base(prs)
        no(sl, s)

        if k == "title":
            put(s, 0.9, 2.05, 10.0, 0.4, sl["eyebrow"].upper(), size=12, color=SUNINK, bold=True, font=MONO)
            put(s, 0.87, 2.5, 12.0, 1.7, sl["big"], size=88, color=INK, bold=True, font=HEAD, line_spacing=0.9)
            put(s, 0.9, 4.35, 10.0, 0.5, sl["sub"], size=22, color=MUTED)
            for j, fact in enumerate(("12 OPEN FEEDS", "5 MODELS TRAINED", "1 VERDICT")):
                chip(s, 0.9 + j * 2.75, 5.35, 2.45, 0.52, fact, fill=CARD, color=MUTED)
            put(s, 0.9, 6.55, 10.0, 0.35, sl["foot"], size=12, color=MUTED, font=MONO)

        elif k == "cards3":
            eyebrow(s, sl["eyebrow"])
            title(s, sl["title"])
            cw, gap, x0 = 3.72, 0.39, 0.9
            for j, (h, body, metric) in enumerate(sl["cards"]):
                x = x0 + j * (cw + gap)
                card(s, x, 2.45, cw, 2.55)
                put(s, x + 0.34, 2.78, cw - 0.68, 0.4, h, size=24, color=SUNINK, bold=True, font=HEAD)
                put(s, x + 0.34, 3.42, cw - 0.68, 0.9, body, size=14.5, color=MUTED, line_spacing=1.3)
                chip(s, x + 0.34, 4.42, 1.55, 0.4, metric.upper(), fill=CARD2, color=MUTED)

        elif k == "rubric":
            eyebrow(s, sl["eyebrow"])
            title(s, sl["title"])
            crit = [("1", "Problem definition", "15%"), ("2", "Originality", "15%"),
                    ("3", "Solution design", "20%"), ("4", "Feasibility", "15%"),
                    ("5", "Demonstrated", "10%"), ("6", "Saudi impact", "15%"),
                    ("7", "Presentation", "10%")]
            for j, (n, nm, wt) in enumerate(crit):
                col, row = j % 4, j // 4
                x = 0.9 + col * 3.05
                y = 2.35 + row * 1.85
                card(s, x, y, 2.75, 1.6)
                chip(s, x + 0.3, y + 0.28, 0.45, 0.34, n, fill=SUN, color=DARKINK, size=12)
                put(s, x + 0.9, y + 0.3, 1.7, 0.32, wt, size=13, color=MUTED, font=MONO)
                put(s, x + 0.3, y + 0.82, 2.2, 0.6, nm, size=15, color=INK, bold=True)
            x8, y8 = 0.9 + 3 * 3.05, 2.35 + 1.85
            card(s, x8, y8, 2.75, 1.6, fill=CARD2)
            put(s, x8 + 0.3, y8 + 0.32, 2.15, 0.5, "5 / 5", size=26, color=SUNINK, bold=True, font=HEAD)
            put(s, x8 + 0.3, y8 + 0.9, 2.15, 0.6, "All seven, evidenced next.", size=13, color=MUTED)
            put(s, 0.9, 6.42, 11.0, 0.4, sl["foot"], size=12, color=MUTED)

        elif k == "criterion":
            put(s, 0.9, 0.6, 8.4, 0.35, f"CRITERION {sl['n']} OF 7   ·   {sl['weight']}",
                size=11, color=SUNINK, bold=True, font=MONO)
            title(s, sl["name"], y=0.98, size=32, w=10.6)
            chip(s, 11.85, 0.95, 0.85, 0.72, "5", fill=SUN, color=DARKINK, size=30, font=HEAD)
            y = 2.25
            for tag, text, src in sl["rows"]:
                card(s, 0.9, y, 11.55, 1.16)
                put(s, 1.2, y + 0.4, 0.3, 0.35, tag, size=15, color=SUNINK, bold=True, font=MONO)
                put(s, 1.62, y + 0.2, 8.95, 0.8, text, size=15, color=INK, line_spacing=1.22,
                    anchor=MSO_ANCHOR.MIDDLE)
                put(s, 10.75, y + 0.26, 1.5, 0.64, src, size=9.5, color=MUTED, font=MONO,
                    align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
                y += 1.33

        elif k == "stats":
            eyebrow(s, sl["eyebrow"])
            title(s, sl["title"])
            for j, (big, lab) in enumerate(sl["stats"]):
                x = 0.9 + j * 3.95
                card(s, x, 2.35, 3.7, 1.95)
                put(s, x + 0.35, 2.6, 3.0, 1.0, big, size=54, color=SUNINK, bold=True, font=HEAD)
                put(s, x + 0.35, 3.66, 3.0, 0.4, lab, size=13, color=MUTED)
            y = 4.72
            for line in sl["rows"]:
                put(s, 0.9, y, 11.4, 0.45, "·  " + line, size=14.5, color=MUTED, line_spacing=1.2)
                y += 0.55

        elif k == "steps":
            eyebrow(s, sl["eyebrow"])
            title(s, sl["title"])
            cw, step = cols(len(sl["steps"]))
            for j, (n, lab) in enumerate(sl["steps"]):
                x = 0.9 + j * step
                card(s, x, 2.3, cw, 1.7)
                chip(s, x + 0.3, 2.58, 0.45, 0.4, n, fill=CARD2, color=SUNINK, size=13)
                put(s, x + 0.3, 3.14, 2.1, 0.7, lab, size=14.5, color=INK, bold=True, line_spacing=1.15)
            y = 4.45
            for line in sl["rows"]:
                card(s, 0.9, y, 11.4, 0.62)
                put(s, 1.2, y + 0.17, 10.8, 0.35, line, size=13.5, color=MUTED)
                y += 0.74

        elif k == "pipeline":
            eyebrow(s, sl["eyebrow"])
            title(s, sl["title"])
            cw, step = cols(len(sl["steps"]))
            for j, (n, lab) in enumerate(sl["steps"]):
                x = 0.9 + j * step
                card(s, x, 2.55, cw, 1.9)
                chip(s, x + 0.3, 2.85, 0.5, 0.5, n, fill=SUN, color=DARKINK, size=16, font=HEAD)
                put(s, x + 0.3, 3.55, 2.1, 0.8, lab, size=15, color=INK, bold=True, line_spacing=1.15)
            put(s, 0.9, 5.1, 11.0, 0.5, sl["note"], size=20, color=SUNINK, bold=True, font=HEAD)

        elif k == "close":
            put(s, 0.9, 2.3, 11.5, 1.4, sl["big"], size=76, color=INK, bold=True, font=HEAD,
                line_spacing=0.95)
            put(s, 0.9, 3.95, 10.5, 0.6, sl["sub"], size=20, color=MUTED)
            put(s, 0.9, 6.5, 10.0, 0.35, sl["foot"], size=12, color=MUTED, font=MONO)

    prs.save(path)
    return len(prs.slides.__iter__.__self__._sldIdLst)


# ───────────────────────────────────────────────────────── html preview
def build_html(path):
    """Same geometry, so the design can be eyeballed in a real browser."""
    def rect(x, y, w, h, bg, *, border=None, r=6):
        b = f"border:1px solid {border};" if border else ""
        return (f'<div style="position:absolute;left:{x*96:.0f}px;top:{y*96:.0f}px;'
                f'width:{w*96:.0f}px;height:{h*96:.0f}px;background:{bg};{b}'
                f'border-radius:{r}px"></div>')

    def txt(x, y, w, h, t, size, color, *, bold=False, mono=False, align="left", ls=1.2):
        fam = "Courier New, monospace" if mono else ("Arial, Helvetica, sans-serif" if bold else
                                                    "Calibri, Carlito, Arial, sans-serif")
        return (f'<div style="position:absolute;left:{x*96:.0f}px;top:{y*96:.0f}px;'
                f'width:{w*96:.0f}px;height:{h*96:.0f}px;color:{color};'
                f'font:{("700 " if bold else "")}{size*1.333:.0f}px/1.2 {fam};'
                f'text-align:{align};line-height:{ls};white-space:pre-wrap">{html.escape(t)}</div>')

    out = ['<html><head><meta charset="utf-8"><style>',
           'body{margin:0;background:#111;font-family:Arial,sans-serif}',
           '.slide{position:relative;width:1280px;height:720px;background:%s;margin:0 0 18px 0;'
           'overflow:hidden}' % PAL["space"],
           '</style></head><body>']
    for i, sl in enumerate(SLIDES):
        k, s = sl["kind"], [f'<div class="slide" id="s{i}">']
        if k == "title":
            s.append(txt(0.9, 2.05, 10, 0.4, sl["eyebrow"].upper(), 12, PAL["sunink"], bold=True, mono=True))
            s.append(txt(0.87, 2.5, 12, 1.7, sl["big"], 88, PAL["ink"], bold=True, ls=0.9))
            s.append(txt(0.9, 4.35, 10, 0.5, sl["sub"], 22, PAL["muted"]))
            for j, fact in enumerate(("12 OPEN FEEDS", "5 MODELS TRAINED", "1 VERDICT")):
                s.append(rect(0.9 + j * 2.75, 5.35, 2.45, 0.52, PAL["card"], border=PAL["line"], r=8))
                s.append(txt(0.9 + j * 2.75, 5.35, 2.45, 0.52, fact, 11, PAL["muted"], mono=True, align="center"))
            s.append(txt(0.9, 6.55, 10, 0.35, sl["foot"], 12, PAL["muted"], mono=True))
        elif k == "cards3":
            s.append(txt(0.9, 0.62, 8, 0.3, sl["eyebrow"].upper(), 11, PAL["sunink"], bold=True, mono=True))
            s.append(txt(0.9, 1.0, 11.5, 0.9, sl["title"], 34, PAL["ink"], bold=True, ls=1.02))
            for j, (h, body, metric) in enumerate(sl["cards"]):
                x = 0.9 + j * 4.11
                s.append(rect(x, 2.45, 3.72, 2.55, PAL["card"], border=PAL["line"], r=14))
                s.append(txt(x + 0.34, 2.78, 3.04, 0.4, h, 24, PAL["sunink"], bold=True))
                s.append(txt(x + 0.34, 3.42, 3.04, 0.9, body, 14.5, PAL["muted"], ls=1.3))
                s.append(rect(x + 0.34, 4.42, 1.55, 0.4, PAL["card2"], r=8))
                s.append(txt(x + 0.34, 4.42, 1.55, 0.4, metric.upper(), 10, PAL["muted"], mono=True, align="center"))
        elif k == "rubric":
            s.append(txt(0.9, 0.62, 8, 0.3, sl["eyebrow"].upper(), 11, PAL["sunink"], bold=True, mono=True))
            s.append(txt(0.9, 1.0, 11.5, 0.9, sl["title"], 34, PAL["ink"], bold=True, ls=1.02))
            crit = [("1", "Problem definition", "15%"), ("2", "Originality", "15%"),
                    ("3", "Solution design", "20%"), ("4", "Feasibility", "15%"),
                    ("5", "Demonstrated", "10%"), ("6", "Saudi impact", "15%"),
                    ("7", "Presentation", "10%")]
            for j, (n, nm, wt) in enumerate(crit):
                col, row = j % 4, j // 4
                x, y = 0.9 + col * 3.05, 2.35 + row * 1.85
                s.append(rect(x, y, 2.75, 1.6, PAL["card"], border=PAL["line"], r=14))
                s.append(rect(x + 0.3, y + 0.28, 0.45, 0.34, PAL["sun"], r=8))
                s.append(txt(x + 0.3, y + 0.28, 0.45, 0.34, n, 12, PAL["darkink"], bold=True, mono=True, align="center"))
                s.append(txt(x + 0.9, y + 0.3, 1.7, 0.32, wt, 13, PAL["muted"], mono=True))
                s.append(txt(x + 0.3, y + 0.82, 2.2, 0.6, nm, 15, PAL["ink"], bold=True))
            x8, y8 = 0.9 + 3 * 3.05, 2.35 + 1.85
            s.append(rect(x8, y8, 2.75, 1.6, PAL["card2"], border=PAL["line"], r=14))
            s.append(txt(x8 + 0.3, y8 + 0.32, 2.15, 0.5, "5 / 5", 26, PAL["sunink"], bold=True))
            s.append(txt(x8 + 0.3, y8 + 0.9, 2.15, 0.6, "All seven, evidenced next.", 13, PAL["muted"]))
            s.append(txt(0.9, 6.42, 11, 0.4, sl["foot"], 12, PAL["muted"]))
        elif k == "criterion":
            s.append(txt(0.9, 0.6, 8.4, 0.35, f"CRITERION {sl['n']} OF 7   ·   {sl['weight']}", 11,
                         PAL["sunink"], bold=True, mono=True))
            s.append(txt(0.9, 0.98, 10.6, 0.9, sl["name"], 32, PAL["ink"], bold=True, ls=1.02))
            s.append(rect(11.85, 0.95, 0.85, 0.72, PAL["sun"], r=16))
            s.append(txt(11.85, 0.95, 0.85, 0.72, "5", 30, PAL["darkink"], bold=True, align="center"))
            y = 2.25
            for tag, text, src in sl["rows"]:
                s.append(rect(0.9, y, 11.55, 1.16, PAL["card"], border=PAL["line"], r=14))
                s.append(txt(1.2, y + 0.4, 0.3, 0.35, tag, 15, PAL["sunink"], bold=True, mono=True))
                s.append(txt(1.62, y + 0.2, 8.95, 0.8, text, 15, PAL["ink"], ls=1.22))
                s.append(txt(10.75, y + 0.26, 1.5, 0.64, src, 9.5, PAL["muted"], mono=True, align="right"))
                y += 1.33
        elif k == "stats":
            s.append(txt(0.9, 0.62, 8, 0.3, sl["eyebrow"].upper(), 11, PAL["sunink"], bold=True, mono=True))
            s.append(txt(0.9, 1.0, 11.5, 0.9, sl["title"], 34, PAL["ink"], bold=True, ls=1.02))
            for j, (big, lab) in enumerate(sl["stats"]):
                x = 0.9 + j * 3.95
                s.append(rect(x, 2.35, 3.7, 1.95, PAL["card"], border=PAL["line"], r=14))
                s.append(txt(x + 0.35, 2.6, 3, 1.0, big, 54, PAL["sunink"], bold=True))
                s.append(txt(x + 0.35, 3.66, 3, 0.4, lab, 13, PAL["muted"]))
            y = 4.72
            for line in sl["rows"]:
                s.append(txt(0.9, y, 11.4, 0.45, "·  " + line, 14.5, PAL["muted"], ls=1.2))
                y += 0.55
        elif k == "steps":
            s.append(txt(0.9, 0.62, 8, 0.3, sl["eyebrow"].upper(), 11, PAL["sunink"], bold=True, mono=True))
            s.append(txt(0.9, 1.0, 11.5, 0.9, sl["title"], 34, PAL["ink"], bold=True, ls=1.02))
            cw, step = cols(len(sl["steps"]))
            for j, (n, lab) in enumerate(sl["steps"]):
                x = 0.9 + j * step
                s.append(rect(x, 2.3, cw, 1.7, PAL["card"], border=PAL["line"], r=14))
                s.append(rect(x + 0.3, 2.58, 0.45, 0.4, PAL["card2"], r=8))
                s.append(txt(x + 0.3, 2.58, 0.45, 0.4, n, 13, PAL["sunink"], bold=True, mono=True, align="center"))
                s.append(txt(x + 0.3, 3.14, 2.1, 0.7, lab, 14.5, PAL["ink"], bold=True))
            y = 4.45
            for line in sl["rows"]:
                s.append(rect(0.9, y, 11.4, 0.62, PAL["card"], border=PAL["line"], r=14))
                s.append(txt(1.2, y + 0.17, 10.8, 0.35, line, 13.5, PAL["muted"]))
                y += 0.74
        elif k == "pipeline":
            s.append(txt(0.9, 0.62, 8, 0.3, sl["eyebrow"].upper(), 11, PAL["sunink"], bold=True, mono=True))
            s.append(txt(0.9, 1.0, 11.5, 0.9, sl["title"], 34, PAL["ink"], bold=True, ls=1.02))
            cw, step = cols(len(sl["steps"]))
            for j, (n, lab) in enumerate(sl["steps"]):
                x = 0.9 + j * step
                s.append(rect(x, 2.55, cw, 1.9, PAL["card"], border=PAL["line"], r=14))
                s.append(rect(x + 0.3, 2.85, 0.5, 0.5, PAL["sun"], r=10))
                s.append(txt(x + 0.3, 2.85, 0.5, 0.5, n, 16, PAL["darkink"], bold=True, align="center"))
                s.append(txt(x + 0.3, 3.55, 2.1, 0.8, lab, 15, PAL["ink"], bold=True))
            s.append(txt(0.9, 5.1, 11, 0.5, sl["note"], 20, PAL["sunink"], bold=True))
        elif k == "close":
            s.append(txt(0.9, 2.3, 11.5, 1.4, sl["big"], 76, PAL["ink"], bold=True, ls=0.95))
            s.append(txt(0.9, 3.95, 10.5, 0.6, sl["sub"], 20, PAL["muted"]))
            s.append(txt(0.9, 6.5, 10, 0.35, sl["foot"], 12, PAL["muted"], mono=True))
        s.append("</div>")
        out.append("".join(s))
    out.append("</body></html>")
    with open(path, "w") as fh:
        fh.write("\n".join(out))
    return len(SLIDES)


if __name__ == "__main__":
    n_pptx = build_pptx(os.path.join(OUT, "SolarGuard-Rubric.pptx"))
    n_html = build_html(os.path.join(OUT, "preview.html"))
    print(f"pptx slides: {n_pptx}   html slides: {n_html}")
    print("wrote", os.path.join(OUT, "SolarGuard-Rubric.pptx"))
    print("wrote", os.path.join(OUT, "preview.html"))
